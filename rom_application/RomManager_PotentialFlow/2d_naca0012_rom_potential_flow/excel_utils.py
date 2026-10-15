import os
import openpyxl

def write_to_data_excel(sheet_name, headers, data_rows, data_file="data.xlsx"):
    """
    Write data to a specific sheet in the unified data.xlsx file.
    If the file exists, load it; if sheet exists, append to it; otherwise create.
    """
    if os.path.exists(data_file):
        wb = openpyxl.load_workbook(data_file)
    else:
        wb = openpyxl.Workbook()
        # Remove default sheet if present
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
    
    # Add headers if sheet doesn't exist
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)
    else:
        ws = wb[sheet_name]
    
    # Append data rows
    for item in data_rows:
        ws.append(item)
    
    wb.save(data_file)
    return data_file
